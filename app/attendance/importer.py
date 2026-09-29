"""Parses a biometric device punch-log export (CSV or XLSX) into
BiometricPunchLog rows, upserts the AttendanceEmployee roster, and
recomputes AttendanceDay for every (emp_code, day) touched.

Expected columns (case-insensitive, order-independent):
Log Date, Device, Emp Code, Name, Company, Department

Log Date arrives as DD-MM-YYYY plus Hr:Min, so it's parsed day-first.
Device carries a raw Entry/Exit marker — stored as-is but not used for any
in/out or duration logic in this pass. Any other columns the export
includes (Category, Status, EmploymentType, ...) are ignored rather than
rejected, so they can be present or absent without affecting the import.
"""
import csv
import io
from dataclasses import dataclass, field
from datetime import datetime

from dateutil import parser as date_parser
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.attendance.models import AttendanceEmployee, BiometricPunchLog, ImportBatch
from app.attendance.rules import recompute_days

REQUIRED_COLUMNS = ["log_date", "emp_code", "name"]
_COLUMN_ALIASES = {
    "log date": "log_date",
    "device": "device",
    "emp code": "emp_code",
    "employee code": "emp_code",
    "name": "name",
    "company": "company",
    "department": "department",
}


@dataclass
class RejectedRow:
    row_number: int
    reason: str


@dataclass
class ImportResult:
    batch_id: str
    row_count: int
    accepted: int
    rejected: list[RejectedRow] = field(default_factory=list)
    employees_affected: int = 0


def _normalize_header(raw_headers: list) -> dict[int, str]:
    """Maps column index -> normalized field name, for whichever of the
    expected columns are present. Unknown columns are ignored rather than
    rejected, so extra device-specific columns don't break the import."""
    mapping: dict[int, str] = {}
    for i, h in enumerate(raw_headers):
        key = str(h or "").strip().lower()
        if key in _COLUMN_ALIASES:
            mapping[i] = _COLUMN_ALIASES[key]
    return mapping


def _rows_from_csv(raw: bytes) -> tuple[list, list[list]]:
    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return [], []
    return rows[0], rows[1:]


def _rows_from_xlsx(raw: bytes) -> tuple[list, list[list]]:
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return [], []
    return rows[0], rows[1:]


def _parse_log_date(value) -> datetime | None:
    """Log Date is DD-MM-YYYY plus Hr:Min — parse day-first so e.g.
    03-04-2026 comes out as 3 April, not 4 March."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    try:
        return date_parser.parse(str(value), dayfirst=True)
    except (ValueError, OverflowError):
        return None


def import_punch_log(db: Session, filename: str, raw: bytes, imported_by: str) -> ImportResult:
    is_xlsx = filename.lower().endswith((".xlsx", ".xlsm"))
    header, data_rows = _rows_from_xlsx(raw) if is_xlsx else _rows_from_csv(raw)
    column_map = _normalize_header(header)

    result = ImportResult(batch_id="", row_count=len(data_rows), accepted=0)
    missing = [c for c in REQUIRED_COLUMNS if c not in column_map.values()]
    if missing:
        result.rejected.append(
            RejectedRow(row_number=0, reason=f"Missing required column(s): {', '.join(missing)}")
        )
        return result

    batch = ImportBatch(filename=filename, imported_by=imported_by, row_count=len(data_rows))
    db.add(batch)
    db.flush()
    result.batch_id = str(batch.id)

    touched: set[tuple[str, "datetime.date"]] = set()
    employee_updates: dict[str, dict] = {}

    for i, raw_row in enumerate(data_rows, start=2):  # spreadsheet row 1 is the header
        if not any(c is not None and str(c).strip() for c in raw_row):
            continue  # skip blank rows

        values = {column_map[idx]: raw_row[idx] for idx in column_map if idx < len(raw_row)}

        log_date = _parse_log_date(values.get("log_date"))
        emp_code = str(values.get("emp_code") or "").strip()
        name = str(values.get("name") or "").strip()

        if log_date is None:
            result.rejected.append(RejectedRow(i, "Unparseable Log Date"))
            continue
        if not emp_code:
            result.rejected.append(RejectedRow(i, "Missing Emp Code"))
            continue

        company = str(values.get("company") or "").strip() or None
        department = str(values.get("department") or "").strip() or None
        device = str(values.get("device") or "").strip() or None

        db.add(
            BiometricPunchLog(
                import_batch_id=batch.id,
                log_date=log_date,
                device=device,
                emp_code=emp_code,
                name=name or emp_code,
                company=company,
                department=department,
            )
        )
        result.accepted += 1
        touched.add((emp_code, log_date.date()))
        employee_updates[emp_code] = {
            "name": name or emp_code,
            "company": company,
            "department": department,
        }

    db.flush()

    for emp_code, fields in employee_updates.items():
        employee = db.get(AttendanceEmployee, emp_code)
        if employee is None:
            employee = AttendanceEmployee(emp_code=emp_code, name=fields["name"])
            db.add(employee)
        employee.name = fields["name"]
        employee.company = fields["company"] or employee.company
        employee.department = fields["department"] or employee.department
        employee.updated_at = datetime.utcnow()

    batch.accepted_count = result.accepted
    batch.rejected_count = len(result.rejected)
    db.commit()

    result.employees_affected = len(employee_updates)
    recompute_days(db, touched)

    return result
